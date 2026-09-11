#!/usr/bin/env python3
"""Convert manuscript-scidata.md to the Scientific Data submission Word file (+ PDF).

Reuses the inline-markup, table and hyperlink helpers from md_to_docx.py. Differences:
citations stay as superscripts (Nature style), figures are embedded at their legends
from scidata/figures/FigureN.png, and the output lands in manuscript-materials/scidata/.

Run:  .venv/bin/python manuscript-materials/md_to_docx_scidata.py
"""
from __future__ import annotations

import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

BASE = Path(__file__).resolve().parent
sys.path.insert(0, str(BASE))
import md_to_docx as m  # noqa: E402

MD_PATH = BASE / "manuscript-scidata.md"
OUT_DIR = BASE / "scidata"
OUT_DOCX = OUT_DIR / "Neurotech_EEG_Dataset_SciData.docx"
FIG_DIR = OUT_DIR / "figures"


def add_figure(doc, fig_num: str) -> bool:
    p = FIG_DIR / f"Figure{fig_num}.png"
    if not p.exists():
        print(f"  WARNING: missing {p}")
        return False
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(p), width=Inches(6.3))
    return True


def process(md_text: str) -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(11)
    lines = md_text.split("\n")
    i, inserted = 0, set()
    while i < len(lines):
        s = lines[i].strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("# ") and not s.startswith("## "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            m.set_run_font(p.add_run(s[2:]), size=16, bold=True)
            i += 1
            continue
        if s.startswith("## "):
            if s[3:].startswith("Figure Legends"):
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            p = doc.add_heading(s[3:], level=1)
            for r in p.runs:
                m.set_run_font(r, size=14, bold=True)
            i += 1
            continue
        if s.startswith("### "):
            p = doc.add_heading(s[4:], level=2)
            for r in p.runs:
                m.set_run_font(r, size=12, bold=True)
            i += 1
            continue
        if s.startswith("|"):
            tl = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                tl.append(lines[i])
                i += 1
            m.add_table_to_doc(doc, m.parse_table(tl))
            doc.add_paragraph()
            continue
        if re.match(r"^\d+\.\s", s):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            m.add_formatted_text(p, s.replace("`", ""))
            i += 1
            continue
        if s.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            m.add_formatted_text(p, s[2:].replace("`", ""))
            i += 1
            continue
        fm = re.match(r"\*\*Figure (\d+)\.\*\*", s)
        if fm:
            n = fm.group(1)
            if n not in inserted and add_figure(doc, n):
                inserted.add(n)
            p = doc.add_paragraph()
            m.add_formatted_text(p, s.replace("`", ""))
            p.paragraph_format.space_after = Pt(14)
            i += 1
            continue
        p = doc.add_paragraph()
        m.add_formatted_text(p, s.replace("`", ""))
        if s.startswith("**Table"):
            p.paragraph_format.space_after = Pt(6)
        i += 1
    return doc


def main():
    OUT_DIR.mkdir(exist_ok=True)
    doc = process(MD_PATH.read_text())
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    doc.save(OUT_DOCX)
    print(f"saved {OUT_DOCX} ({OUT_DOCX.stat().st_size/1024:.0f} KB)")
    r = subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(OUT_DIR), str(OUT_DOCX)],
                       capture_output=True, text=True, timeout=300)
    pdf = OUT_DOCX.with_suffix(".pdf")
    print("pdf:", pdf if pdf.exists() else f"FAILED {r.stderr[:300]}")


if __name__ == "__main__":
    main()
